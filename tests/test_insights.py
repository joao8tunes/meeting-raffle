import io
from datetime import datetime, time

import pytest
from openpyxl import load_workbook

from raffle import charts, samples
from raffle.draw import Entry, Round
from raffle.formats import engagement_kind
from raffle.insights import ScoreWeights, build_insights, pseudonyms
from raffle.loader import load_file, load_people_list
from raffle.models import KIND_LIST, KIND_REPORT, Attendance, Meeting
from raffle.people import Rules, build_roster
from raffle.reports import (events_frame, insights_workbook, matrix_frame, overview_frame, people_frame,
                            prizes_frame)

DAY = datetime(2026, 3, 3)


def at(hour, minute=0, day=DAY):
    return day.replace(hour=hour, minute=minute)


def utf16(text: str) -> bytes:
    import codecs

    return codecs.BOM_UTF16_LE + text.replace("\n", "\r\n").encode("utf-16-le")


@pytest.fixture(scope="module")
def series():
    meetings = []
    for name, data in samples.demo_files():
        meetings += load_file(data, name).meetings
    return meetings


# ---------------------------------------------------------------------------------------------- engagement parsing

@pytest.mark.parametrize("label, kind", [
    ("Engagement: Reaction - Applause", "reactions"), ("Envolvimento: Reação - Aplauso", "reactions"),
    ("Compromisso: Câmera Ligada", "camera"), ("Turned camera on", "camera"),
    ("Compromisso: Levantar as mãos", "raised_hands"), ("Mãos levantadas", "raised_hands"),
    ("Compromisso: Desativar Mudo", "unmutes"), ("Mudo desativado", "unmutes"), ("Stummschaltung aufheben", "unmutes"),
    ("Chat messages", "chat"), ("Duração da Reunião", None), ("Email", None), ("Tipo de Compromisso", None),
])
def test_engagement_kind(label, kind):
    assert engagement_kind(label) == kind


def test_teams_report_engagement_columns_and_log_are_not_double_counted():
    report = utf16(
        "1. Resumo\nTítulo da reunião\tWeekly sync\n"
        "Hora de início\t7/16/26, 4:00:00 PM\nHora de término\t7/16/26, 5:00:00 PM\n"
        "\n2. Participantes\n"
        "Nome\tPrimeira Entrada\tÚltima Saída\tDuração da Reunião\tEmail\tFunção\t"
        "Envolvimento: Reação - Aplauso\tCompromisso: Câmera Ligada\tCompromisso: Desativar Mudo\n"
        "Lee, Kim\t7/16/26, 4:00:00 PM\t7/16/26, 5:00:00 PM\t1h\tkim@example.com\tApresentador\t2\t\t1\n"
        "Ana Nova\t7/16/26, 4:10:00 PM\t7/16/26, 4:50:00 PM\t40m\tana@example.com\tApresentador\t\t\t\n"
        "\n4. Compromisso da Reunião\nNome\tTipo de Compromisso\tTempo\n"
        "Lee, Kim\tReação enviada \"aplausos\"\t7/16/26, 4:10:00 PM\n"
        "Lee, Kim\tReação enviada \"aplausos\"\t7/16/26, 4:20:00 PM\n"
        "Lee, Kim\tMudo desativado\t7/16/26, 4:30:00 PM\n"
        "Lee, Kim\tAtivou a câmera\t7/16/26, 4:31:00 PM\n"
    )
    meeting = load_file(report, "report.csv").meetings[0]
    assert meeting.kind == KIND_REPORT and meeting.engagement_tracked
    people = {a.name: a for a in meeting.attendees}
    assert people["Lee, Kim"].engagement == {"reactions": 2, "unmutes": 1, "camera": 1}
    assert len(people["Lee, Kim"].engagement_events) == 4
    assert people["Ana Nova"].engagement == {}


def test_sample_series_has_engagement_except_the_classic_list(series):
    tracked = [m.engagement_tracked for m in series]
    assert tracked.count(True) == 3 and tracked.count(False) == 1


def test_no_show_rows_are_not_counted_as_prizes():
    history = load_people_list(samples.previous_winners_csv(), "winners.csv", keep_repeats=True)
    assert len(history) == 3 and len(set(history)) == 2  # one person won twice, the no-show row is skipped
    assert len(load_people_list(samples.previous_winners_csv(), "winners.csv")) == 2


# ------------------------------------------------------------------------------------------------------ insights

def test_event_metrics_new_returning_peak_and_retention():
    first = Meeting("a.csv", KIND_REPORT, [
        Attendance("Ana", "ana@example.com", sessions=[(at(17), at(18))]),
        Attendance("Bo", "bo@example.com", sessions=[(at(17), at(17, 30))]),
    ], title="First", start=at(17), end=at(18), first_seen=at(17), last_seen=at(18))
    later = DAY.replace(day=10)
    second = Meeting("b.csv", KIND_REPORT, [
        Attendance("Ana", "ana@example.com", sessions=[(at(17, day=later), at(18, day=later))]),
        Attendance("Cy", "cy@example.com", sessions=[(at(17, 15, later), at(17, 45, later))]),
    ], title="Second", start=at(17, day=later), end=at(18, day=later))
    meetings = [second, first]  # upload order must not matter
    insights = build_insights(meetings, build_roster(meetings))

    one, two = insights.events
    assert (one.title, two.title) == ("First", "Second")
    assert (one.audience, one.new_people, one.returning_people) == (2, 2, 0)
    assert (two.audience, two.new_people, two.returning_people) == (2, 1, 1)
    assert one.peak == 2 and one.peak_at == at(17)
    assert one.avg_minutes == pytest.approx(45)
    assert one.avg_share == pytest.approx(0.75)
    assert one.stayed_share == pytest.approx(0.5)
    assert dict(one.timeline)[40] == 1

    ana = next(p for p in insights.people if p.name == "Ana")
    assert (ana.events, ana.streak, ana.attendance_rate, round(ana.minutes)) == (2, 2, 1.0, 120)
    assert insights.overview.returning_share == pytest.approx(1 / 3)
    assert insights.distribution == {1: 2, 2: 1}


def test_time_window_limits_audience_and_minutes():
    meeting = Meeting("a.csv", KIND_REPORT, [
        Attendance("Setup crew", sessions=[(at(16), at(16, 50))]),
        Attendance("Audience", sessions=[(at(17), at(18))]),
    ], start=at(16), end=at(18), first_seen=at(16), last_seen=at(18))
    window = (time(17), time(18))
    insights = build_insights([meeting], build_roster([meeting], Rules(window=window)), window)
    event = insights.events[0]
    assert event.audience == 1 and event.duration_minutes == pytest.approx(60)
    assert "time window" in insights.notes[0]


def test_relevance_score_uses_available_components_and_weights():
    big = Meeting("a", KIND_REPORT, [Attendance(f"P{i}", sessions=[(at(17), at(17, 30))]) for i in range(4)],
                  start=at(17), end=at(18), engagement_tracked=True)
    big.attendees[0].engagement = {"reactions": 3}
    small_day = DAY.replace(day=5)
    small = Meeting("b", KIND_REPORT, [Attendance("P0", sessions=[(at(17, day=small_day), at(18, day=small_day))])],
                    start=at(17, day=small_day), end=at(18, day=small_day))
    roster = build_roster([big, small])
    first, second = build_insights([big, small], roster).events
    assert first.score_parts == {"audience": 1.0, "retention": 0.5, "engagement": 0.25}
    assert first.score == pytest.approx(100 * (0.4 * 1 + 0.4 * 0.5 + 0.2 * 0.25))
    assert "engagement" not in second.score_parts
    assert second.score == pytest.approx(100 * (0.4 * 0.25 + 0.4 * 1.0) / 0.8)
    only_audience = build_insights([big, small], roster, weights=ScoreWeights(1, 0, 0)).events
    assert [round(e.score) for e in only_audience] == [100, 25]


def test_prizes_from_session_and_history(series):
    roster = build_roster(series)
    person = roster.eligible[0]
    draw = Round(1, "Mug", 1, False, [], [Entry(person.key, person.name, person.email)])
    history = [person.email, person.name.upper(), "nobody@example.com"]
    insights = build_insights(series, roster, rounds=[draw], prize_history=history)
    winner = next(p for p in insights.people if p.key == person.key)
    assert (winner.prizes_session, winner.prizes_history) == (1, 2)
    assert insights.unmatched_prizes == 1
    assert insights.overview.repeat_winners == 1
    no_show = Round(2, "Mug", 1, False, [], [Entry(person.key, person.name)], no_shows=[person.key])
    assert build_insights(series, roster, rounds=[no_show]).overview.prizes == 0


def test_plain_lists_count_for_audience_without_minutes():
    meeting = Meeting("list.txt", KIND_LIST, [Attendance("Walk-in")])
    insights = build_insights([meeting], build_roster([meeting]))
    assert insights.events[0].audience == 1
    assert insights.events[0].duration_minutes is None and insights.events[0].timeline == []
    assert any("without join/leave times" in note for note in insights.notes)


def test_pseudonyms_are_stable_and_do_not_follow_names():
    keys = [f"name:{letter}" for letter in "abcdefghij"]
    codes = pseudonyms(keys)
    assert codes == pseudonyms(reversed(keys))
    assert sorted(codes.values()) == [f"Person {i:03d}" for i in range(1, 11)]
    assert [codes[k] for k in keys] != sorted(codes.values())


# ------------------------------------------------------------------------------------------------ reports & charts

def test_report_tables_and_anonymized_workbook(series):
    roster = build_roster(series)
    history = load_people_list(samples.previous_winners_csv(), "w.csv", keep_repeats=True)
    insights = build_insights(series, roster, prize_history=history)
    assert len(events_frame(insights)) == 4
    assert "Events" in overview_frame(insights)["Metric"].tolist()
    assert set(matrix_frame(insights).columns) >= {e.label for e in insights.events}
    assert not prizes_frame(insights).empty

    anonymous = people_frame(insights, anonymize=True)
    assert "Email" not in anonymous.columns
    assert anonymous["Person"].str.match(r"^Person \d{3}$").all()

    workbook = load_workbook(io.BytesIO(insights_workbook(insights, series, anonymize=True)))
    assert {"Overview", "Events", "People", "Minutes per event", "Audience over time", "Engagement", "Prizes",
            "Files"} <= set(workbook.sheetnames)
    assert workbook["Events"]._charts and workbook["Engagement"]._charts
    text = " ".join(str(cell.value) for sheet in workbook for row in sheet.iter_rows() for cell in row
                    if cell.value is not None)
    for person in samples.sample_people():
        assert person.full not in text and (not person.email or person.email not in text)


@pytest.mark.parametrize("dark", [False, True])
def test_charts_build_valid_specs(series, dark):
    insights = build_insights(series, build_roster(series))
    theme = charts.Theme(dark)
    specs = [
        charts.audience_per_event(insights, theme), charts.minutes_per_event(insights, theme),
        charts.events_attended(insights, theme), charts.relevance(insights, theme),
        charts.audience_over_time(insights, charts.default_overlay(insights), theme),
        charts.engagement_per_event(insights, theme),
        charts.engagement_over_time(next(e for e in insights.events if e.engagement_timeline), theme),
    ]
    for spec in specs:
        assert spec.to_dict()["$schema"]
