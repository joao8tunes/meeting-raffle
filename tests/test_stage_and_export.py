import io
import json
import re
import shutil
import subprocess

import pandas as pd
import pytest
from openpyxl import load_workbook

from raffle.draw import Entry, run_round
from raffle.export import participants_frame, results_workbook, to_csv, winners_frame
from raffle.models import KIND_LIST, Attendance, Meeting
from raffle.people import build_roster
from raffle.stage import (ELIMINATION, LOTTERY, MAX_ON_STAGE, MAX_WINNERS, MODE_HELP, MODES, RACE, RACE_THEMES,
                          REVEAL, SPOTLIGHT, WHEEL, build_idle_stage, build_stage, effective_mode, lottery_tickets)

EVIL = '</script><img src=x onerror="alert(1)">'


def payload(html: str) -> dict:
    match = re.search(r"const DATA = (.*?);\n", html)
    return json.loads(match.group(1))


def test_every_show_is_described_and_limited():
    assert set(MODE_HELP) == set(MODES)
    assert set(MAX_WINNERS) == set(MODES) - {REVEAL}
    assert len(RACE_THEMES) >= 7


def test_names_cannot_break_out_of_the_stage_script():
    pool = [Entry("k1", EVIL), Entry("k2", "=cmd|' /C calc'!A0")] + [Entry(f"k{i}", f"P{i}") for i in range(3, 90)]
    draw = run_round(1, "<b>Prize</b>", pool, 2)
    for mode in MODES:
        html = build_stage(draw, mode)
        assert "</script><img" not in html
        assert html.count("</script>") == 1
        data = payload(html)
        on_stage = data["names"]
        assert all(on_stage[w["slot"]] == w["name"] for w in data["winners"])
        assert [w["name"] for w in data["winners"]] == [w.name for w in draw.winners]


@pytest.mark.parametrize("mode", [WHEEL, RACE, SPOTLIGHT, ELIMINATION])
def test_everyone_is_on_stage_for_realistic_audiences(mode):
    pool = [Entry(f"k{i:04d}", f"Person {i}") for i in range(1500)]
    data = payload(build_stage(run_round(1, "", pool, 3), mode))
    assert len(data["names"]) == 1500
    assert "All 1,500 people in the draw are in the show" in data["caption"]


def test_huge_audiences_show_a_sample_that_always_includes_the_winners():
    pool = [Entry(f"k{i:05d}", f"Person {i}") for i in range(MAX_ON_STAGE + 500)]
    draw = run_round(1, "", pool, 5)
    data = payload(build_stage(draw, RACE))
    assert len(data["names"]) == MAX_ON_STAGE
    assert all(data["names"][w["slot"]] == w["name"] for w in data["winners"])
    assert f"all {MAX_ON_STAGE + 500:,} people" in data["caption"]


def test_lottery_tickets_follow_alphabetical_order_and_reach_the_winner():
    pool = [Entry(f"k{i:04d}", name) for i, name in enumerate(["Zoë Prism", "andré Marble", "Blake Quasar"]
                                                                  + [f"Person {i:04d}" for i in range(247)])]
    draw = run_round(1, "", pool, 2)
    data = payload(build_stage(draw, LOTTERY))
    ordered = lottery_tickets(draw.pool)
    assert ordered[0].name == "andré Marble" and ordered[-1].name == "Zoë Prism"
    lottery = data["lottery"]
    assert (lottery["tickets"], lottery["digits"]) == (250, 3)
    assert sorted(data["names"]) == sorted(winner.name for winner in draw.winners)  # no need to send everyone
    for winner, ticket in zip(draw.winners, lottery["draws"]):
        assert ordered[ticket["ticket"] - 1].key == winner.key
        numbers = [number for number, _ in ticket["neighbors"]]
        assert ticket["ticket"] in numbers and len(numbers) <= 10
        assert all(number // 10 == ticket["ticket"] // 10 for number in numbers)


def test_small_lottery_lists_every_ticket():
    draw = run_round(1, "", [Entry(f"k{i}", f"P{i}") for i in range(4)], 1)
    lottery = payload(build_stage(draw, LOTTERY))["lottery"]
    assert lottery["digits"] == 1 and len(lottery["draws"][0]["neighbors"]) == 4


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is not installed")
def test_stage_script_is_valid_javascript(tmp_path):
    html = build_stage(run_round(1, "", [Entry("k", "Kim")], 1), WHEEL)
    script = re.search(r"<script>(.*)</script>", html, re.S).group(1)
    path = tmp_path / "stage.js"
    path.write_text(script, encoding="utf-8")
    result = subprocess.run(["node", "--check", str(path)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_emails_only_on_stage_when_requested():
    draw = run_round(1, "", [Entry("k", "Kim", "kim@example.com")], 1)
    assert "kim@example.com" not in build_stage(draw, REVEAL)
    assert "kim@example.com" in build_stage(draw, REVEAL, show_email=True)


def test_big_rounds_fall_back_to_instant_reveal():
    assert effective_mode(WHEEL, 6)[0] == REVEAL
    assert effective_mode(RACE, 3) == (RACE, "")
    assert "Ready when you are" in payload(build_idle_stage("Mug", 10))["idleTitle"]


def test_exports_neutralize_spreadsheet_formulas():
    pool = [Entry("k1", "=HYPERLINK(\"http://x\")", "a@example.com"), Entry("k2", "+55 11 99999", "")]
    draw = run_round(1, "@prize", pool, 2)
    csv_text = to_csv(winners_frame([draw])).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and "'+55" in csv_text and "'@prize" in csv_text

    meeting = Meeting("list.csv", KIND_LIST, [Attendance("=SUM(A1)"), Attendance("Normal Name")])
    roster = build_roster([meeting])
    workbook = load_workbook(io.BytesIO(results_workbook([draw], roster, [meeting])))
    assert set(workbook.sheetnames) == {"Winners", "Audit", "Pools", "Participants", "Files"}
    values = [cell.value for row in workbook["Participants"].iter_rows() for cell in row]
    assert "'=SUM(A1)" in values
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet.iter_rows() for cell in row)


def test_participants_frame_can_hide_emails():
    meeting = Meeting("list.csv", KIND_LIST, [Attendance("Kim", "kim@example.com")])
    roster = build_roster([meeting])
    assert "Email" in participants_frame(roster, [meeting]).columns
    assert "Email" not in participants_frame(roster, [meeting], include_email=False).columns
    assert isinstance(winners_frame([], include_email=False), pd.DataFrame)
