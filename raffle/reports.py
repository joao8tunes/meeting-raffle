"""Insight tables and the downloadable Excel report (with native charts and optional anonymization)."""

from __future__ import annotations

import io
from typing import Dict, List, Optional, Sequence

import pandas as pd

from .export import _sanitize, files_frame
from .insights import Insights, pseudonyms
from .models import ENGAGEMENT_KINDS, ENGAGEMENT_LABELS, Meeting

# Categorical slots (validated for color-vision deficiencies) used by the Excel charts.
SERIES_COLORS = ["2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948"]


def _pct(value: Optional[float]) -> Optional[float]:
    return None if value is None else round(100 * value, 1)


def _round(value: Optional[float], digits: int = 0) -> Optional[float]:
    if value is None:
        return None
    return int(round(value)) if digits == 0 else round(value, digits)


def _names(insights: Insights, anonymize: bool) -> Dict[str, str]:
    if anonymize:
        return pseudonyms(person.key for person in insights.people)
    return {person.key: person.name for person in insights.people}


def _kinds_used(insights: Insights) -> List[str]:
    used = {kind for event in insights.events for kind in event.engagement}
    return [kind for kind in ENGAGEMENT_KINDS if kind in used]


def overview_frame(insights: Insights) -> pd.DataFrame:
    o = insights.overview

    def pct(value: Optional[float]) -> str:
        return "–" if value is None else f"{100 * value:.0f}%"

    def num(value: Optional[float], digits: int = 1) -> str:
        return "–" if value is None else f"{value:,.{digits}f}"

    rows = [
        ("Events", f"{o.events}"),
        ("Unique people", f"{o.people}"),
        ("Attendances (people × events)", f"{o.attendances}"),
        ("Person-hours", num(o.person_hours)),
        ("Average audience per event", num(o.avg_audience)),
        ("Average minutes per attendance", num(o.avg_minutes)),
        ("Average events per person", num(o.avg_events_per_person)),
        ("People who came back (2+ events)", pct(o.returning_share)),
        ("Stayed until the end", pct(o.stayed_share)),
        ("Engaged (reacted, spoke, raised hand or camera on)", pct(o.engaged_share)),
        ("Prizes awarded", f"{o.prizes}"),
        ("People awarded", f"{o.people_awarded}"),
        ("People awarded more than once", f"{o.repeat_winners}"),
        ("Events attended by winners (average)", num(o.winners_avg_events)),
    ]
    frame = pd.DataFrame(rows, columns=["Metric", "Value"])
    notes = pd.DataFrame([("Note", note) for note in insights.notes], columns=["Metric", "Value"])
    return pd.concat([frame, notes], ignore_index=True) if len(notes) else frame


def events_frame(insights: Insights, anonymize: bool = False) -> pd.DataFrame:
    kinds = _kinds_used(insights)
    rows = []
    for e in insights.events:
        row = {
            "#": e.number,
            "Event": e.label,
            "Date": e.day.isoformat() if e.day else None,
            "Start": e.start.strftime("%H:%M") if e.start and e.duration_minutes is not None else None,
            "End": e.end.strftime("%H:%M") if e.end and e.duration_minutes is not None else None,
            "Duration (min)": _round(e.duration_minutes),
            "Audience": e.audience,
            "New": e.new_people,
            "Returning": e.returning_people,
            "Peak": e.peak,
            "Peak at": e.peak_at.strftime("%H:%M") if e.peak_at else None,
            "Avg minutes": _round(e.avg_minutes),
            "Median minutes": _round(e.median_minutes),
            "Avg attended (%)": _pct(e.avg_share),
            "Stayed to the end (%)": _pct(e.stayed_share),
            "Engaged (%)": _pct(e.engaged_share),
        }
        for kind in kinds:
            row[ENGAGEMENT_LABELS[kind]] = e.engagement.get(kind, 0) if e.engagement_tracked else None
        row["Relevance"] = _round(e.score)
        if not anonymize:
            row["File"] = e.source
        row["Format"] = e.kind
        rows.append(row)
    return pd.DataFrame(rows)


def people_frame(insights: Insights, anonymize: bool = False, include_email: bool = True) -> pd.DataFrame:
    names = _names(insights, anonymize)
    kinds = _kinds_used(insights)
    rows = []
    for p in insights.people:
        row = {"Person": names[p.key]}
        if include_email and not anonymize:
            row["Email"] = p.email
        row.update({
            "Roles": ", ".join(p.roles),
            "Events": p.events,
            "Attendance rate (%)": _pct(p.attendance_rate),
            "Total minutes": _round(p.minutes),
            "Avg minutes per event": _round(p.avg_minutes),
            "Stayed to the end": p.stayed,
            "Longest streak": p.streak,
            "First event": p.first_event,
            "Last event": p.last_event,
        })
        for kind in kinds:
            row[ENGAGEMENT_LABELS[kind]] = p.engagement.get(kind, 0)
        row.update({
            "Prizes (this session)": p.prizes_session,
            "Prizes (before)": p.prizes_history,
            "Prizes": p.prizes,
            "Eligible now": p.eligible,
        })
        rows.append(row)
    frame = pd.DataFrame(rows)
    if not frame.empty:
        frame = frame.sort_values(["Events", "Total minutes", "Person"], ascending=[False, False, True],
                                  na_position="last").reset_index(drop=True)
    return frame


def matrix_frame(insights: Insights, anonymize: bool = False) -> pd.DataFrame:
    """One row per person, one column per event: minutes attended ("attended" for lists without times)."""
    names = _names(insights, anonymize)
    rows = []
    for p in insights.people:
        row = {"Person": names[p.key], "Events": p.events}
        for e in insights.events:
            if e.number not in p.minutes_by_event:
                row[e.label] = None
            else:
                minutes = p.minutes_by_event[e.number]
                row[e.label] = _round(minutes) if minutes is not None else "attended"
        rows.append(row)
    frame = pd.DataFrame(rows)
    if not frame.empty:
        frame = frame.sort_values(["Events", "Person"], ascending=[False, True]).reset_index(drop=True)
    return frame


def timeline_frame(insights: Insights) -> pd.DataFrame:
    rows = []
    for e in insights.events:
        for minute, people in e.timeline:
            rows.append({
                "Event": e.label,
                "Minute": round(minute, 1),
                "People connected": people,
                "Share of audience (%)": _pct(people / e.audience) if e.audience else None,
            })
    return pd.DataFrame(rows, columns=["Event", "Minute", "People connected", "Share of audience (%)"])


def engagement_frame(insights: Insights) -> pd.DataFrame:
    kinds = _kinds_used(insights)
    rows = []
    for e in insights.events:
        if not e.engagement_tracked:
            continue
        row = {"Event": e.label, "Audience": e.audience, "Engaged people": e.engaged_people,
               "Engaged (%)": _pct(e.engaged_share)}
        for kind in kinds:
            row[ENGAGEMENT_LABELS[kind]] = e.engagement.get(kind, 0)
        rows.append(row)
    return pd.DataFrame(rows)


def prizes_frame(insights: Insights, anonymize: bool = False, include_email: bool = True) -> pd.DataFrame:
    names = _names(insights, anonymize)
    rows = []
    for p in sorted(insights.people, key=lambda p: (-p.prizes, -p.events, names[p.key])):
        if not p.prizes:
            continue
        row = {"Person": names[p.key]}
        if include_email and not anonymize:
            row["Email"] = p.email
        row.update({
            "Prizes": p.prizes,
            "This session": p.prizes_session,
            "Before": p.prizes_history,
            "Events attended": p.events,
            "Total minutes": _round(p.minutes),
            "Eligible now": p.eligible,
        })
        rows.append(row)
    return pd.DataFrame(rows)


def _wide_timeline(insights: Insights) -> pd.DataFrame:
    frame = timeline_frame(insights)
    if frame.empty:
        return frame
    frame["Minute"] = frame["Minute"].round().astype(int)
    wide = frame.pivot_table(index="Minute", columns="Event", values="People connected", aggfunc="max")
    ordered = [e.label for e in insights.events if e.label in wide.columns]
    return wide[ordered].reset_index()


def insights_workbook(insights: Insights, meetings: Sequence[Meeting], anonymize: bool = False) -> bytes:
    from openpyxl.chart import BarChart, LineChart, Reference

    sheets: Dict[str, pd.DataFrame] = {
        "Overview": overview_frame(insights),
        "Events": events_frame(insights, anonymize),
        "People": people_frame(insights, anonymize),
        "Minutes per event": matrix_frame(insights, anonymize),
    }
    wide = _wide_timeline(insights)
    if not wide.empty:
        sheets["Audience over time"] = wide
    engagement = engagement_frame(insights)
    if not engagement.empty:
        sheets["Engagement"] = engagement
    prizes = prizes_frame(insights, anonymize)
    if not prizes.empty:
        sheets["Prizes"] = prizes
    files = files_frame(meetings)
    if anonymize:
        files = files.drop(columns=["File"])
    sheets["Files"] = files

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        for name, frame in sheets.items():
            _sanitize(frame).to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "B2"
            for cells in sheet.columns:
                width = max((len(str(cell.value)) for cell in cells if cell.value is not None), default=8)
                sheet.column_dimensions[cells[0].column_letter].width = min(max(10, width + 2), 48)

        def column(frame: pd.DataFrame, title: str) -> int:
            return list(frame.columns).index(title) + 1

        def paint(chart, colors: Sequence[str]) -> None:
            for series, color in zip(chart.series, colors):
                series.graphicalProperties.solidFill = color
                series.graphicalProperties.line.solidFill = color

        events = sheets["Events"]
        count = len(events)
        if count:
            sheet = writer.sheets["Events"]
            categories = Reference(sheet, min_col=column(events, "Event"), min_row=2, max_row=count + 1)
            audience = BarChart()
            audience.type, audience.grouping, audience.overlap = "col", "stacked", 100
            audience.title, audience.y_axis.title = "Audience per event", "People"
            audience.add_data(Reference(sheet, min_col=column(events, "New"), max_col=column(events, "Returning"),
                                        min_row=1, max_row=count + 1), titles_from_data=True)
            audience.set_categories(categories)
            paint(audience, SERIES_COLORS[1::-1])  # New = orange, Returning = blue
            audience.width, audience.height = 18, 8
            sheet.add_chart(audience, f"B{count + 4}")

            minutes = BarChart()
            minutes.type, minutes.title, minutes.y_axis.title = "col", "Average minutes per event", "Minutes"
            minutes.add_data(Reference(sheet, min_col=column(events, "Avg minutes"), min_row=1, max_row=count + 1),
                             titles_from_data=True)
            minutes.set_categories(categories)
            minutes.legend = None
            paint(minutes, SERIES_COLORS[:1])
            minutes.width, minutes.height = 18, 8
            sheet.add_chart(minutes, f"L{count + 4}")

        if "Audience over time" in sheets and len(wide.columns) > 1:
            sheet = writer.sheets["Audience over time"]
            line = LineChart()
            line.title, line.y_axis.title, line.x_axis.title = "People connected", "People", "Minutes since start"
            series_count = min(len(wide.columns) - 1, len(SERIES_COLORS))
            line.add_data(Reference(sheet, min_col=2, max_col=1 + series_count, min_row=1, max_row=len(wide) + 1),
                          titles_from_data=True)
            line.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=len(wide) + 1))
            for series, color in zip(line.series, SERIES_COLORS):
                series.graphicalProperties.line.solidFill = color
                series.graphicalProperties.line.width = 20000
                series.smooth = False
            line.width, line.height = 22, 10
            sheet.add_chart(line, f"{_letter(len(wide.columns) + 2)}2")

        if "Engagement" in sheets:
            kinds = [ENGAGEMENT_LABELS[k] for k in _kinds_used(insights)]
            if kinds:
                sheet = writer.sheets["Engagement"]
                rows = len(engagement)
                chart = BarChart()
                chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
                chart.title, chart.y_axis.title = "Engagement per event", "Count"
                chart.add_data(Reference(sheet, min_col=column(engagement, kinds[0]),
                                         max_col=column(engagement, kinds[-1]), min_row=1, max_row=rows + 1),
                               titles_from_data=True)
                chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=rows + 1))
                paint(chart, SERIES_COLORS)
                chart.width, chart.height = 18, 8
                sheet.add_chart(chart, f"B{rows + 4}")
    return buffer.getvalue()


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)
