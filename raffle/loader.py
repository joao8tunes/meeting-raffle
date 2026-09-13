"""Read uploaded files (CSV/TSV/TXT/XLSX) into meetings, whatever their encoding or delimiter."""

from __future__ import annotations

import codecs
import csv
import io
import re
import zipfile
from datetime import datetime, time, timedelta
from pathlib import PurePath
from typing import List, Optional, Tuple

from .formats import interpret_rows, is_header_label
from .models import KIND_LIST, Attendance, LoadResult, Meeting
from .names import email_key, name_key
from .text import clean_cell
from .timeparse import AUTO

TEXT_EXTENSIONS = ("csv", "tsv", "txt")
EXCEL_EXTENSIONS = ("xlsx", "xlsm")
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS + EXCEL_EXTENSIONS

_DELIMITERS = ("\t", ";", ",", "|")
_LIST_PREFIX_RE = re.compile(r"^\s*(?:[-*•·]|\d{1,4}\s*[.)-])\s+")
_MAX_BYTES = 50 * 1024 * 1024


def decode_text(data: bytes) -> str:
    """Decode bytes from exports made by different apps and operating systems."""
    for bom, encoding in ((codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"),
                          (codecs.BOM_UTF16_BE, "utf-16")):
        if data.startswith(bom):
            return data.decode(encoding, errors="replace")

    sample = data[:4096]
    if len(sample) >= 4:
        if sample[1::2].count(0) > len(sample) // 4:
            return data.decode("utf-16-le", errors="replace")
        if sample[0::2].count(0) > len(sample) // 4:
            return data.decode("utf-16-be", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    try:
        return data.decode("cp1252")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _looks_binary(data: bytes) -> bool:
    sample = data[:2048]
    return b"\x00" in sample and not (sample.startswith((codecs.BOM_UTF16_LE, codecs.BOM_UTF16_BE))
                                      or sample[1::2].count(0) > len(sample) // 4
                                      or sample[0::2].count(0) > len(sample) // 4)


def split_rows(lines: List[str], delimiter: str) -> List[List[str]]:
    rows = []
    for line in lines:
        # Parse line by line so a stray quote cannot swallow the rest of the file.
        cells = next(csv.reader([line], delimiter=delimiter), []) if line.strip() else []
        rows.append([clean_cell(cell) for cell in cells])
    while rows and not any(rows[-1]):
        rows.pop()
    return rows


def _delimiter_candidates(lines: List[str]) -> List[str]:
    non_empty = [line for line in lines if line.strip()][:200]
    if not non_empty:
        return []
    scores = []
    for delimiter in _DELIMITERS:
        share = sum(1 for line in non_empty if delimiter in line) / len(non_empty)
        if share > 0:
            scores.append((share, -_DELIMITERS.index(delimiter), delimiter))
    return [delimiter for _, _, delimiter in sorted(scores, reverse=True)]


def plain_list(lines: List[str], source: str) -> Optional[Meeting]:
    """One person per line. Commas stay inside names ("Doe, Jane")."""
    attendees: dict = {}
    for line in lines:
        text = line
        for delimiter in ("\t", ";"):
            if delimiter in text:
                text = next((cell for cell in text.split(delimiter) if cell.strip()), "")
                break
        text = clean_cell(_LIST_PREFIX_RE.sub("", text))
        if not text or not any(ch.isalpha() for ch in text) or is_header_label(text):
            continue
        email = email_key(text)
        key = f"email:{email}" if email else f"name:{name_key(text)}"
        attendees.setdefault(key, Attendance(name=text, email=email))
    if not attendees:
        return None
    return Meeting(source=source, kind=KIND_LIST, attendees=list(attendees.values()))


Interpretation = Tuple[List[Meeting], int]


def _as_interpretation(meeting: Optional[Meeting]) -> Interpretation:
    return ([meeting], 1) if meeting else ([], 0)


def _best_interpretation(candidates: List[Interpretation]) -> List[Meeting]:
    """Prefer the most confident reading of a file, then the one that finds the most people."""
    scored = [(level, sum(len(m.attendees) for m in meetings), -position, meetings)
              for position, (meetings, level) in enumerate(candidates) if meetings]
    return max(scored, key=lambda item: item[:3])[3] if scored else []


def load_text(data: bytes, source: str, date_order: str = AUTO) -> List[Meeting]:
    lines = decode_text(data).splitlines()
    candidates: List[Interpretation] = []
    for delimiter in _delimiter_candidates(lines):
        candidates.append(interpret_rows(split_rows(lines, delimiter), source, date_order))
        if candidates[-1][1] == 3:
            break
    candidates.append(_as_interpretation(plain_list(lines, source)))
    return _best_interpretation(candidates)


def _excel_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        total = int(value.total_seconds())
        return f"{total // 3600}:{total % 3600 // 60:02d}:{total % 60:02d}"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return clean_cell(value)


def load_excel(data: bytes, source: str, date_order: str = AUTO) -> List[Meeting]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    meetings = []
    try:
        sheets = workbook.worksheets
        for sheet in sheets:
            rows = [[_excel_cell(value) for value in row] for row in sheet.iter_rows(values_only=True)]
            while rows and not any(rows[-1]):
                rows.pop()
            if not rows:
                continue
            label = source if len(sheets) == 1 else f"{source} › {sheet.title}"
            lines = ["\t".join(cell.replace("\t", " ") for cell in row) for row in rows]
            meetings.extend(_best_interpretation([
                interpret_rows(rows, label, date_order),
                _as_interpretation(plain_list(lines, label)),
            ]))
    finally:
        workbook.close()
    return meetings


def load_file(data: bytes, source: str, date_order: str = AUTO) -> LoadResult:
    """Load one file. Never raises: problems are reported in ``LoadResult.error``."""
    result = LoadResult(source=source)
    extension = PurePath(source).suffix.lower().lstrip(".")
    try:
        if not data:
            result.error = "The file is empty."
        elif len(data) > _MAX_BYTES:
            result.error = "The file is too large (limit: 50 MB)."
        elif extension == "xls":
            result.error = "Old Excel files (.xls) are not supported. Save it as .xlsx or .csv and try again."
        elif extension in EXCEL_EXTENSIONS or data.startswith(b"PK\x03\x04"):
            result.meetings = load_excel(data, source, date_order)
        elif _looks_binary(data):
            result.error = "This does not look like a CSV, TXT or Excel file."
        else:
            result.meetings = load_text(data, source, date_order)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        result.error = f"Could not read this file ({exc.__class__.__name__}). Is it a valid spreadsheet?"
    except Exception as exc:  # noqa: BLE001 - one bad upload must never break the app
        result.error = f"Unexpected problem while reading this file ({exc.__class__.__name__})."

    if not result.error and not result.meetings:
        result.error = "No participant names were found in this file."
    return result


def load_people_list(data: bytes, source: str, keep_repeats: bool = False) -> List[str]:
    """Names and emails from any supported file (exclusion lists, prize history).

    With ``keep_repeats`` a person listed in several rows (e.g. two prizes) is returned once per row.
    """
    result = load_file(data, source)
    # An email is more precise than a name: never exclude a namesake when the email is known.
    return [attendee.email or attendee.name
            for meeting in result.meetings for attendee in meeting.attendees
            for _ in range(attendee.mentions if keep_repeats else 1)]
